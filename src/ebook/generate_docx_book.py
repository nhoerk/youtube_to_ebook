from pathlib import Path
from src.core.paths import output_dir
from docx import Document
import os
import re


def _output_name(extension):
    title = os.getenv(
        "YOUTUBE_TO_EBOOK_BOOK_TITLE",
        "Membentengi Akidah, Memurnikan Tauhid",
    )
    slug = re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")
    return f"{slug}.{extension}"


def generate_docx():

    manuscript = output_dir() / "manuscript.md"

    text = manuscript.read_text(
        encoding="utf-8"
    )

    doc = Document()

    for line in text.splitlines():

        line = line.strip()

        if not line:
            continue

        if line.startswith("# "):
            doc.add_heading(
                line.replace("# ", ""),
                level=1,
            )

        elif line.startswith("## "):
            doc.add_heading(
                line.replace("## ", ""),
                level=2,
            )

        else:
            doc.add_paragraph(line)

    output_file = output_dir() / _output_name("docx")

    doc.save(output_file)

    print(
        f"DOCX tersimpan: {output_file}"
    )

    return output_file